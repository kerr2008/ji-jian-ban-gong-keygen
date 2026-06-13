# -*- coding: utf-8 -*-
"""
数据工具模块
文本批量替换 / 编码转换 / 手机号邮箱格式化
"""

import os
import re
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from datetime import datetime
from pathlib import Path
import threading


def batch_replace(files, rules, output_dir, progress_callback=None):
    """批量文本替换
    rules: [{"old":"", "new":"", "regex":False}, ...]
    """
    total = len(files)
    results = []
    for i, fp in enumerate(files):
        ext = Path(fp).suffix.lower()
        changes = 0
        try:
            text = _read_file(fp)
            if text is None:
                results.append({"文件": os.path.basename(fp), "状态": "读取失败"})
                continue

            for rule in rules:
                old = rule.get("old", "")
                new = rule.get("new", "")
                is_regex = rule.get("regex", False)
                if not old:
                    continue
                if is_regex:
                    new_text, count = re.subn(old, new, text)
                else:
                    new_text, count = text.replace(old, new), text.count(old)
                text = new_text
                changes += count

            if changes > 0:
                out_name = Path(fp).stem + "_替换" + ext
                out_path = os.path.join(output_dir, out_name)
                _write_file(out_path, text, ext)
                results.append({"文件": os.path.basename(fp), "状态": "替换" + str(changes) + "处"})
            else:
                results.append({"文件": os.path.basename(fp), "状态": "无匹配"})

        except Exception as e:
            results.append({"文件": os.path.basename(fp), "状态": "错误: " + str(e)})

        if progress_callback:
            progress_callback(i + 1, total)
    return results


def _read_file(fp):
    ext = Path(fp).suffix.lower()
    if ext in (".txt", ".csv", ".html", ".xml", ".json", ".md"):
        for enc in ["utf-8", "gbk", "gb18030"]:
            try:
                with open(fp, "r", encoding=enc) as f:
                    return f.read()
            except:
                continue
    elif ext in (".docx", ".doc"):
        from docx import Document
        doc = Document(fp)
        return "\n".join(p.text for p in doc.paragraphs)
    return None


def _write_file(fp, text, ext):
    if ext in (".txt", ".csv", ".html", ".xml", ".json", ".md"):
        with open(fp, "w", encoding="utf-8") as f:
            f.write(text)
    elif ext in (".docx", ".doc"):
        from docx import Document
        doc = Document()
        for line in text.split("\n"):
            doc.add_paragraph(line)
        doc.save(fp)


def batch_convert_encoding(files, output_dir, target_enc="utf-8", progress_callback=None):
    """批量编码转换"""
    total = len(files)
    results = []
    for i, fp in enumerate(files):
        try:
            # 探测原编码
            src_enc = "utf-8"
            for enc in ["utf-8", "gbk", "gb18030", "utf-16"]:
                try:
                    with open(fp, "r", encoding=enc) as f:
                        text = f.read()
                    src_enc = enc
                    break
                except:
                    continue

            out_name = Path(fp).stem + "_utf8" + Path(fp).suffix
            out_path = os.path.join(output_dir, out_name)
            with open(out_path, "w", encoding=target_enc) as f:
                f.write(text)
            results.append({"文件": os.path.basename(fp), "源编码": src_enc, "目标编码": target_enc})

        except Exception as e:
            results.append({"文件": os.path.basename(fp), "状态": "错误: " + str(e)})

        if progress_callback:
            progress_callback(i + 1, total)
    return results


def format_phone_email(text, output_path, progress_callback=None):
    """提取并格式化手机号和邮箱"""
    phone_pat = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")
    email_pat = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")

    phones = list(dict.fromkeys(phone_pat.findall(text)))
    emails = list(dict.fromkeys(email_pat.findall(text)))

    # 格式化手机号
    formatted_phones = []
    for p in phones:
        fp = p[:3] + "-" + p[3:7] + "-" + p[7:]
        formatted_phones.append(fp)

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("提取结果\n")
        f.write("=" * 30 + "\n\n")
        f.write("手机号 ({}个):\n".format(len(phones)))
        for p in formatted_phones:
            f.write("  " + p + "\n")
        f.write("\n邮箱 ({}个):\n".format(len(emails)))
        for e in emails:
            f.write("  " + e + "\n")
        f.write("\n原始数据:\n")
        f.write("=" * 30 + "\n")
        f.write(text)

    return phones, emails


class DataToolFrame(ttk.Frame):
    def __init__(self, parent, theme):
        super().__init__(parent)
        self.theme = theme
        self.selected_files = []
        self.progress_var = tk.DoubleVar()
        self.status_var = tk.StringVar(value="准备就绪")
        self.last_output = None
        self._build_ui()

    def _build_ui(self):
        c = self.theme.colors
        tk.Label(self, text="数据工具 — 文本替换 / 编码转换 / 格式化",
                 font=("Microsoft YaHei", 11, "bold"),
                 bg=c["panel_bg"], fg=c["fg"], anchor="w").pack(fill="x", pady=(0, 15))

        mf = tk.Frame(self, bg=c["panel_bg"])
        mf.pack(fill="x", pady=(0, 10))
        tk.Label(mf, text="操作：", bg=c["panel_bg"], fg=c["fg"],
                 font=("Microsoft YaHei", 9)).pack(side="left")
        self.mode = tk.StringVar(value="replace")
        for text, val in [("批量替换", "replace"), ("编码转换", "encoding"), ("格式化", "format")]:
            tk.Radiobutton(mf, text=text, variable=self.mode, value=val,
                           bg=c["panel_bg"], fg=c["fg"], selectcolor=c["input_bg"],
                           font=("Microsoft YaHei", 9),
                           activebackground=c["panel_bg"], activeforeground=c["fg"],
                           command=self._on_mode_change).pack(side="left", padx=(8, 0))

        ff = tk.Frame(self, bg=c["panel_bg"])
        ff.pack(fill="x", pady=(0, 5))
        tk.Label(ff, text="选择文件：", bg=c["panel_bg"], fg=c["fg"],
                 font=("Microsoft YaHei", 9)).pack(side="left")
        self.file_label = tk.Label(ff, text="未选择", bg=c["panel_bg"], fg=c["disabled_fg"],
                                    font=("Microsoft YaHei", 9))
        self.file_label.pack(side="left", fill="x", expand=True, padx=(10, 0))
        self.btn_select = tk.Button(ff, text="📁 选择文件",
                                    command=self._select_files,
                                    bg=c["accent"], fg="white",
                                    font=("Microsoft YaHei", 9),
                                    relief="flat", padx=12, pady=3, cursor="hand2")
        self.btn_select.pack(side="right")

        self.param_frame = tk.LabelFrame(self, text="参数设置", bg=c["panel_bg"], fg=c["fg"],
                                          font=("Microsoft YaHei", 9), padx=10, pady=5)
        self.param_frame.pack(fill="x", pady=(0, 10))
        self.param_content = tk.Frame(self.param_frame, bg=c["panel_bg"])
        self.param_content.pack(fill="x")
        self._build_replace_params()

        af = tk.Frame(self, bg=c["panel_bg"])
        af.pack(fill="x", pady=(5, 10))
        self.btn_start = tk.Button(af, text="🚀 开始处理",
                                   command=self._start_process,
                                   bg=c["success"], fg="white",
                                   font=("Microsoft YaHei", 10, "bold"),
                                   relief="flat", padx=20, pady=5,
                                   cursor="hand2", state="disabled")
        self.btn_start.pack(side="left", padx=(0, 10))
        self.btn_open = tk.Button(af, text="📂 打开输出文件夹",
                                  command=lambda: self._open_last(),
                                  bg=c["sidebar_bg"], fg=c["fg"],
                                  font=("Microsoft YaHei", 9),
                                  relief="flat", padx=12, pady=3,
                                  cursor="hand2", state="disabled")
        self.btn_open.pack(side="left")

        pf = tk.Frame(self, bg=c["panel_bg"])
        pf.pack(fill="x", pady=(5, 5))
        self.progress = ttk.Progressbar(pf, variable=self.progress_var, maximum=100, mode="determinate")
        self.progress.pack(fill="x", pady=(0, 3))
        self.status_label = tk.Label(pf, textvariable=self.status_var,
                                     bg=c["panel_bg"], fg=c["disabled_fg"], font=("Microsoft YaHei", 8))
        self.status_label.pack(anchor="w")

        rf = tk.LabelFrame(self, text="执行结果", bg=c["panel_bg"], fg=c["fg"],
                            font=("Microsoft YaHei", 9), padx=5, pady=5)
        rf.pack(fill="both", expand=True)
        self.result_text = tk.Text(rf, bg=c["input_bg"], fg=c["fg"],
                                    font=("Consolas", 9), wrap="word",
                                    relief="flat", padx=5, pady=5, state="disabled")
        self.result_text.pack(fill="both", expand=True)

    def _build_replace_params(self):
        self._clear_params()
        c = self.theme.colors
        f = tk.Frame(self.param_content, bg=c["panel_bg"])
        f.pack(fill="x", pady=2)
        tk.Label(f, text="替换规则：", bg=c["panel_bg"], fg=c["fg"],
                 font=("Microsoft YaHei", 9)).pack(side="left")
        tk.Label(f, text="查找：", bg=c["panel_bg"], fg=c["fg"],
                 font=("Microsoft YaHei", 9)).pack(side="left", padx=(10, 0))
        self.replace_old = tk.Entry(f, bg=c["input_bg"], fg=c["fg"],
                                     font=("Microsoft YaHei", 9), relief="flat", bd=1, width=20)
        self.replace_old.pack(side="left", padx=(5, 5))
        tk.Label(f, text="替换为：", bg=c["panel_bg"], fg=c["fg"],
                 font=("Microsoft YaHei", 9)).pack(side="left")
        self.replace_new = tk.Entry(f, bg=c["input_bg"], fg=c["fg"],
                                     font=("Microsoft YaHei", 9), relief="flat", bd=1, width=20)
        self.replace_new.pack(side="left", padx=(5, 0))
        self.use_regex = tk.BooleanVar(value=False)
        tk.Checkbutton(f, text="正则", variable=self.use_regex,
                       bg=c["panel_bg"], fg=c["fg"], selectcolor=c["input_bg"],
                       font=("Microsoft YaHei", 9),
                       activebackground=c["panel_bg"]).pack(side="left", padx=(10, 0))

    def _build_encoding_params(self):
        self._clear_params()
        c = self.theme.colors
        tk.Label(self.param_content, text="目标编码：", bg=c["panel_bg"], fg=c["fg"],
                 font=("Microsoft YaHei", 9)).pack(side="left")
        self.target_enc = tk.StringVar(value="utf-8")
        ttk.Combobox(self.param_content, textvariable=self.target_enc,
                      values=["utf-8", "gbk", "gb18030", "utf-16"],
                      state="readonly", width=10).pack(side="left", padx=(5, 0))
        tk.Label(self.param_content, text="（将文件转为统一编码，解决乱码）",
                 bg=c["panel_bg"], fg=c["disabled_fg"], font=("Microsoft YaHei", 8)).pack(side="left", padx=(10, 0))

    def _build_format_params(self):
        self._clear_params()
        c = self.theme.colors
        tk.Label(self.param_content, text="从文本中提取手机号（格式化分段）和邮箱地址",
                 bg=c["panel_bg"], fg=c["disabled_fg"], font=("Microsoft YaHei", 9)).pack(side="left")

    def _clear_params(self):
        for w in self.param_content.winfo_children():
            w.destroy()

    def _on_mode_change(self):
        {
            "replace": self._build_replace_params,
            "encoding": self._build_encoding_params,
            "format": self._build_format_params,
        }.get(self.mode.get(), lambda: None)()

    def _select_files(self):
        files = filedialog.askopenfilenames(
            title="选择文件",
            filetypes=[("文本文件", "*.txt *.csv *.md *.html *.xml *.json"), ("Word文件", "*.docx"), ("所有文件", "*.*")]
        )
        if files:
            self.selected_files = files
            self.file_label.config(text=str(len(files)) + " 个文件", fg=self.theme.get("fg"))
            self.btn_start.config(state="normal")

    def _start_process(self):
        if not self.selected_files:
            messagebox.showwarning("提示", "请先选择文件")
            return
        self.btn_start.config(state="disabled", text="⏳ 处理中...")
        self.progress_var.set(0)
        self.result_text.config(state="normal")
        self.result_text.delete("1.0", "end")
        self.result_text.config(state="disabled")
        thread = threading.Thread(target=self._process, daemon=True)
        thread.start()

    def _process(self):
        output_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "极简办公_输出")
        os.makedirs(output_dir, exist_ok=True)

        def log(msg):
            self.after(0, lambda: self._append_log(msg))

        def cb(c, t):
            self.after(0, lambda: self.progress_var.set(c / t * 100))

        try:
            mode = self.mode.get()
            if mode == "replace":
                rules = [{"old": self.replace_old.get(), "new": self.replace_new.get(),
                          "regex": self.use_regex.get()}]
                results = batch_replace(self.selected_files, rules, output_dir, cb)
                log("✅ 替换完成！")
                for r in results:
                    log("  {}: {}".format(r["文件"], r["状态"]))

            elif mode == "encoding":
                enc = self.target_enc.get()
                results = batch_convert_encoding(self.selected_files, output_dir, enc, cb)
                log("✅ 编码转换完成！")
                for r in results:
                    log("  {}: {} → {}".format(r["文件"], r["源编码"], r["目标编码"]))

            elif mode == "format":
                # 合并所有文本
                all_text = ""
                for fp in self.selected_files:
                    t = _read_file(fp)
                    if t:
                        all_text += t + "\n"
                out = os.path.join(output_dir, "格式化结果_" + datetime.now().strftime("%Y%m%d_%H%M%S") + ".txt")
                phones, emails = format_phone_email(all_text, out)
                log("✅ 格式化完成！")
                log("  提取手机号: {} 个".format(len(phones)))
                for p in phones:
                    log("    " + p[:3] + "-" + p[3:7] + "-" + p[7:])
                log("  提取邮箱: {} 个".format(len(emails)))
                for e in emails:
                    log("    " + e)
                self.last_output = out

            self.after(0, lambda: self.btn_open.config(state="normal"))

        except Exception as e:
            log("❌ 错误: " + str(e))

        self.progress_var.set(100)
        self.after(0, lambda: self.btn_start.config(state="normal", text="🚀 开始处理"))
        self.after(0, lambda: self.status_var.set("处理完成 ✅"))

    def _append_log(self, text):
        self.result_text.config(state="normal")
        self.result_text.insert("end", text + "\n")
        self.result_text.see("end")
        self.result_text.config(state="disabled")

    def _open_last(self):
        if self.last_output:
            p = self.last_output
            if os.path.isfile(p):
                os.startfile(os.path.dirname(p))
            else:
                os.startfile(p)

    def apply_theme(self):
        try:
            self.result_text.config(bg=self.theme.colors["input_bg"], fg=self.theme.colors["fg"])
        except:
            pass